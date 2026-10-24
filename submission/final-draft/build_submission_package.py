#!/usr/bin/env python3
"""Build editable KEIT report drafts and clearly labeled visual QA previews.

The PPTX and DOCX start from the official templates. PDF files generated here
are rasterized layout previews from the same content; they are not native
PowerPoint/Word exports and must not be presented as such.
"""
from __future__ import annotations

from copy import deepcopy
from pathlib import Path
import hashlib
import json
import textwrap

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pptx import Presentation
from pptx.dml.color import RGBColor as PptColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches as PptInches, Pt as PptPt
from PIL import Image, ImageDraw, ImageFont


ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent
EVIDENCE = OUT / "evidence"
PREVIEW = EVIDENCE / "rendered"
NOTICE = ROOT / "artifacts/experiment_results.json"
PPT_TEMPLATE = ROOT / "submission/forms/draft/성과보고서_양식_PPT_원본복사.pptx"
DOCX_TEMPLATE = ROOT / "submission/forms/draft/성과보고서_양식_Word_원본복사.docx"
E2E_DIR = OUT / "evidence/source"

PPT_OUT = OUT / "VillageCoverage_성과보고서_20261023.pptx"
DOCX_OUT = OUT / "VillageCoverage_성과보고서_20261023.docx"
PPT_PDF = OUT / "VillageCoverage_성과보고서_20261023_PPTX_검증미리보기.pdf"
DOCX_PDF = OUT / "VillageCoverage_성과보고서_20261023_DOCX_검증미리보기.pdf"

GREEN = "176B4A"
DARK = "18382D"
MINT = "EAF4EF"
PALE = "F5F8F6"
GOLD = "A66A11"
GOLD_BG = "FFF5E4"
RED = "A83D34"
GRAY = "52615B"
LIGHT = "D8E2DC"
WHITE = "FFFFFF"

CORE = "요청이 적다는 이유로 생활서비스가 필요한 마을이 공급계획에서 사라지지 않도록 돕는 AI 기반 의사결정 지원 시스템."
TECH_NAME = "VillageCoverage — 농촌 생활서비스 공급계획 AI 의사결정 지원 시스템"
IDEA_NAME = "요청 기록의 공백을 드러내고 생활서비스 공급안을 비교하는 마을 단위 의사결정 지원"
SIM_LABEL = "SIMULATED FOR PRE-R&D"
SOURCE_SHA = "b6ff52adb2fa8b4287e6c62250802d8c07196907"
EXPECTED_EXPERIMENT_SHA = "4d984db66fc9b90624dbbb1a26245b6e7f52b3b89e77e44a381c05dfb2cf24eb"


def load_metrics():
    raw = NOTICE.read_bytes()
    actual_sha = hashlib.sha256(raw).hexdigest()
    if actual_sha != EXPECTED_EXPERIMENT_SHA:
        raise ValueError(f"experiment artifact SHA changed: {actual_sha}; re-audit before rebuilding")
    doc = json.loads(raw)
    ex = doc["experiment_1_request_count_vs_balanced"]
    return doc, ex, actual_sha


def prep_evidence():
    EVIDENCE.mkdir(parents=True, exist_ok=True)
    PREVIEW.mkdir(parents=True, exist_ok=True)
    dash = Image.open(E2E_DIR / "dashboard.png").convert("RGB")
    comp = Image.open(E2E_DIR / "scenario-comparison.png").convert("RGB")
    repl = Image.open(E2E_DIR / "plan-replanned.png").convert("RGB")
    # These are direct crops from the local E2E screenshots; no UI or data is
    # redrawn. The separate panels remain labeled by their capture route.
    dash.crop((180, 0, 1350, 930)).save(EVIDENCE / "Dashboard_local-E2E.png")
    dash.crop((205, 1180, 1345, 2240)).save(EVIDENCE / "Dashboard_map_local-E2E.png")
    comp.crop((180, 0, 1350, 950)).save(EVIDENCE / "ScenarioCompare_local-E2E.png")
    repl.crop((190, 0, 1340, 1180)).save(EVIDENCE / "Replan_local-E2E.png")


def ppt_rgb(value: str) -> PptColor:
    return PptColor.from_string(value)


def add_text(slide, text, x, y, w, h, *, size=18, color=DARK, bold=False,
             font="Pretendard", align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE,
             margin=0.04, fit=False):
    box = slide.shapes.add_textbox(PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = PptInches(margin)
    tf.margin_top = tf.margin_bottom = PptInches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_after = PptPt(0)
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = PptPt(size)
    run.font.bold = bold
    run.font.color.rgb = ppt_rgb(color)
    if fit:
        from pptx.enum.text import MSO_AUTO_SIZE
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return box


def add_rect(slide, x, y, w, h, fill, *, line=None, radius=False):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    sh = slide.shapes.add_shape(kind, PptInches(x), PptInches(y), PptInches(w), PptInches(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = ppt_rgb(fill)
    sh.line.color.rgb = ppt_rgb(line or fill)
    sh.line.width = PptPt(1)
    return sh


def add_label(slide, text, x, y, w, *, fill=MINT, color=GREEN, size=12):
    add_rect(slide, x, y, w, 0.34, fill, line=fill, radius=True)
    add_text(slide, text, x + 0.07, y + 0.01, w - 0.14, 0.31,
             size=size, color=color, bold=True, align=PP_ALIGN.CENTER, margin=0)


def replace_title(slide, exact):
    for shape in slide.shapes:
        if shape.name == "TextBox 2":
            original_run = shape.text_frame.paragraphs[0].runs[0]
            color = original_run.font.color.rgb
            shape.text_frame.clear()
            paragraph = shape.text_frame.paragraphs[0]
            run = paragraph.add_run()
            run.text = exact
            run.font.name = "Pretendard SemiBold"
            run.font.size = PptPt(33.33)
            run.font.bold = False
            if color:
                run.font.color.rgb = color
            shape.width = PptInches(19.6)
            return


def build_ppt():
    data, ex, artifact_sha = load_metrics()
    baseline, balanced = ex["request_count_baseline"], ex["balanced"]
    total_areas = len(ex["request_count_service_units_by_area"])
    prs = Presentation(str(PPT_TEMPLATE))
    # The template has exactly four official pages and the original logo,
    # heading rules, page numbers, and instruction notes remain in place.
    slides = prs.slides
    replace_title(slides[0], "1. AI 기술명")
    for slide in slides:
        for shape in slide.shapes:
            if shape.left + shape.width > prs.slide_width:
                # The official template header line protrudes by 0.02 in.
                # Trim only that edge so the exported page cannot clip it.
                shape.width = prs.slide_width - shape.left

    # Slide 1: real local application captures, cropped only from actual E2E
    # screenshots. Captions explicitly distinguish local test from Render.
    s = slides[0]
    add_text(s, TECH_NAME, 1.35, 3.10, 23.7, 0.55, size=24, bold=True, color=DARK)
    add_text(s, CORE, 1.39, 3.63, 23.55, 0.63, size=17, color=GRAY)
    add_text(s, "네 정책안 비교 · 같은 지역/예산 조건", 1.42, 4.28, 11.20, 0.36, size=13, color=GREEN, bold=True)
    add_text(s, "마을별 계획 지도 · 조사 필요/미배정 표시", 13.88, 4.28, 11.20, 0.36, size=13, color=GREEN, bold=True)
    s.shapes.add_picture(str(EVIDENCE / "ScenarioCompare_local-E2E.png"), PptInches(1.42), PptInches(4.68), width=PptInches(11.30), height=PptInches(9.17))
    s.shapes.add_picture(str(EVIDENCE / "Dashboard_map_local-E2E.png"), PptInches(14.58), PptInches(4.68), width=PptInches(9.85), height=PptInches(9.17))
    add_text(s, "로컬 Playwright E2E 캡처 · 합성 데모 데이터 · 2026-10-13 · 소스 b6ff52a · 원격 Render 화면 아님",
             1.45, 13.94, 23.7, 0.30, size=10, color=GRAY, valign=MSO_ANCHOR.TOP)

    # Slide 2: preserve the official title and instruction sentence; item/award
    # identifiers remain unresolved rather than being fabricated.
    s = slides[1]
    add_rect(s, 1.38, 3.14, 23.7, 0.62, GOLD_BG, line=GOLD_BG, radius=True)
    add_text(s, "이전 아이디어 번호·정확한 선정 상태·5점 가점 자격: USER_CONFIRMATION_REQUIRED",
             1.60, 3.22, 23.2, 0.42, size=15, color=GOLD, bold=True)
    flow = [
        ("01  문제", "요청 기록이 적은 마을을\n필요가 없는 곳으로\n오판할 위험"),
        ("02  근거 구조화", "AI 보조로 비정형 기록을\n검토 초안으로 정리.\n부족·충돌은 조사 필요"),
        ("03  제약 최적화", "CP-SAT이 예산·공급·이동·\n일정 조건 아래 네 정책안을\n계산해 비교"),
        ("04  담당자 검토", "결과·가정·미배정 사유를\n검토하고 정책을 선택.\n자동 승인하지 않음"),
    ]
    start, card_w, gap = 1.40, 5.65, 0.45
    for i, (head, body) in enumerate(flow):
        x = start + i * (card_w + gap)
        add_rect(s, x, 4.12, card_w, 2.60, PALE, line=LIGHT, radius=True)
        add_label(s, head, x + 0.20, 4.35, card_w - 0.40, fill=MINT, color=GREEN, size=14)
        add_text(s, body, x + 0.25, 4.92, card_w - 0.50, 1.35, size=17, color=DARK, valign=MSO_ANCHOR.TOP)
        if i < 3:
            add_text(s, "→", x + card_w + 0.03, 5.02, 0.39, 0.40, size=21, color=GREEN, bold=True, align=PP_ALIGN.CENTER)
    add_text(s, "시제품 정책 프리셋", 1.43, 7.08, 7.5, 0.42, size=18, bold=True)
    policies = [
        ("효율 중심", "같은 예산에서 서비스 회차 효율"),
        ("균형", "회차·권역 분산·조사 필요 고려"),
        ("소외 최소화", "서비스 공백이 긴 권역 우선"),
        ("최소보장", "권역별 최소 회차 조건 우선"),
    ]
    for i, (head, body) in enumerate(policies):
        x = 1.42 + i * 6.08
        add_rect(s, x, 7.60, 5.70, 1.75, WHITE, line=LIGHT, radius=True)
        add_text(s, head, x + 0.18, 7.81, 5.34, 0.38, size=16, bold=True, color=GREEN)
        add_text(s, body, x + 0.18, 8.27, 5.34, 0.78, size=13, color=GRAY, valign=MSO_ANCHOR.TOP)
    add_rect(s, 1.42, 9.78, 23.70, 1.15, MINT, line=MINT, radius=True)
    add_text(s, "AI와 최적화의 역할", 1.68, 9.91, 3.0, 0.35, size=16, bold=True, color=GREEN)
    add_text(s, "AI는 근거를 정리하는 보조 경로, CP-SAT은 명시된 목적·제약을 계산하는 solver입니다. 정책 가중치와 단가는 시제품 선택이며 현장 합의값이 아닙니다.",
             4.55, 9.89, 19.7, 0.62, size=15, color=DARK)
    add_text(s, "사용자 측 사전질의 답변은 원 제안 취지를 설명한 자료이며, KEIT의 개별 구현 승인으로 해석하지 않음.",
             1.50, 11.15, 23.45, 0.40, size=12, color=GRAY)

    # Slide 3: executable features and one exact synthetic experiment. The
    # comparison is explicitly only request-count baseline vs balanced.
    s = slides[2]
    add_text(s, "구현 범위와 검증 경계", 1.43, 3.13, 7.2, 0.45, size=20, bold=True, color=DARK)
    features = [
        ("근거·불확실성", "관측 없음/부족/충돌을 저수요로 확정하지 않고 조사 필요·UNKNOWN으로 남김."),
        ("AI 기록 구조화", "Gemini 설정 시 구조화 출력, 미설정 시 결정적 fallback. 본 E2E는 외부 모델 키 없이 실행."),
        ("네 정책 계산", "OR-Tools CP-SAT으로 예산·공급·이동·일정 제약 아래 효율·균형·최소보장·소외 최소화 계산."),
        ("불참·검토 산출", "합성 제공자 불참 후 새 계획 버전과 Decision Memo/PDF·CSV 흐름을 로컬 E2E에서 확인."),
    ]
    for i, (head, body) in enumerate(features):
        y = 3.72 + i * 1.47
        add_rect(s, 1.42, y, 10.72, 1.28, PALE, line=LIGHT, radius=True)
        add_text(s, head, 1.66, y + 0.13, 2.30, 0.36, size=15, bold=True, color=GREEN)
        add_text(s, body, 4.06, y + 0.12, 7.78, 0.95, size=13, color=DARK, valign=MSO_ANCHOR.TOP)

    add_rect(s, 12.53, 3.13, 12.58, 6.75, WHITE, line=LIGHT, radius=True)
    add_text(s, "통제 합성 실험 · 두 정책만 비교", 12.82, 3.36, 11.98, 0.48, size=18, bold=True, color=DARK)
    add_label(s, SIM_LABEL, 12.85, 3.93, 3.38, fill=GOLD_BG, color=GOLD, size=11)
    add_text(s, f"{data['region']} · {total_areas}개 권역 · 예산 {ex['budget_won']:,}원 · seed {data['seed']} · 조사 필요=합성 관측 4건 미만",
             12.83, 4.43, 11.95, 0.55, size=12, color=GRAY)
    headers = ["지표", "요청량 기준", "균형"]
    rows = [
        ("서비스 단위", str(baseline["served_units"]), str(balanced["served_units"])),
        ("배정 권역", f"{baseline['covered_areas']} / {total_areas}", f"{balanced['covered_areas']} / {total_areas}"),
        ("미배정 권역", str(total_areas - baseline["covered_areas"]), str(total_areas - balanced["covered_areas"])),
        ("조사 필요 권역 배정", f"{baseline['survey_required_covered']} / {baseline['survey_required_areas']}", f"{balanced['survey_required_covered']} / {balanced['survey_required_areas']}"),
        ("모델 이동비", f"₩{baseline['travel_cost_won']:,}", f"₩{balanced['travel_cost_won']:,}"),
    ]
    x0, y0, widths, row_h = 12.86, 5.13, [4.4, 3.55, 3.45], 0.63
    cx = x0
    for i, h in enumerate(headers):
        add_rect(s, cx, y0, widths[i], row_h, MINT, line=MINT)
        add_text(s, h, cx + 0.08, y0 + 0.06, widths[i] - 0.16, row_h - 0.12,
                 size=12, color=GREEN, bold=True, align=PP_ALIGN.CENTER if i else PP_ALIGN.LEFT)
        cx += widths[i]
    for r, row in enumerate(rows):
        cx = x0
        for c, val in enumerate(row):
            fill = WHITE if r % 2 else PALE
            add_rect(s, cx, y0 + row_h * (r + 1), widths[c], row_h, fill, line=LIGHT)
            add_text(s, val, cx + 0.08, y0 + row_h * (r + 1) + 0.05, widths[c] - 0.16, row_h - 0.10,
                     size=12, color=DARK, bold=(c > 0), align=PP_ALIGN.CENTER if c else PP_ALIGN.LEFT)
            cx += widths[c]
    add_text(s, "모델 이동비는 공개 고정 도로행렬과 시제품 단가에 따른 추정치로, 실제 지자체 집행비가 아닙니다.",
             12.88, 9.02, 11.90, 0.52, size=11, color=GRAY)
    add_rect(s, 1.43, 10.18, 23.67, 1.12, GOLD_BG, line=GOLD_BG, radius=True)
    add_text(s, "해석 한계", 1.69, 10.31, 1.80, 0.35, size=15, color=GOLD, bold=True)
    add_text(s, "합성 수요·관측·공급·가격으로 계산한 Pre-R&D 예시입니다. 실제 주민 서비스, 9개 마을 추가 지원, 정책 우열, 사회적 효과를 입증하지 않습니다.",
             3.50, 10.25, 21.00, 0.57, size=14, color=DARK)
    add_text(s, f"출처: artifacts/experiment_results.json · SHA-256 {artifact_sha} · 실험 항목 experiment_1_request_count_vs_balanced · UI E2E 결과와 별도",
             1.48, 11.48, 23.65, 0.35, size=10, color=GRAY)

    # Slide 4: distinguish engineering constraints, mitigations and unproven
    # expected value. Do not turn these expectations into measured impact.
    s = slides[3]
    add_text(s, "개발 애로점", 1.43, 3.13, 11.25, 0.45, size=20, bold=True, color=DARK)
    add_text(s, "설계에서 취한 대응", 13.18, 3.13, 11.55, 0.45, size=20, bold=True, color=DARK)
    issues = [
        ("현장 주민 수요·미기록 필요 자료 부족", "관측 부재를 0으로 대체하지 않고, 부족·충돌은 조사 필요/UNKNOWN으로 유지"),
        ("제공자 가격·용량·참여가 미확정", "합성 운영값을 실제 계약조건과 분리하고 불참 뒤 후보/비용 재계산 경로 제공"),
        ("이동·방문일정·solver 복잡도", "입력 제약·solver 상태를 함께 표시; FEASIBLE은 최적성 증명으로 간주하지 않음"),
        ("데이터 출처·개인정보·행정 보안", "근거 출처와 공개/합성 구분, 사람 검토 단계를 유지; 실기관 보안 검증은 향후 과제"),
    ]
    for i, (left, right) in enumerate(issues):
        y = 3.72 + i * 1.27
        fill = PALE if i % 2 == 0 else WHITE
        add_rect(s, 1.43, y, 11.15, 1.07, fill, line=LIGHT, radius=True)
        add_text(s, left, 1.68, y + 0.12, 10.66, 0.78, size=15, color=DARK, valign=MSO_ANCHOR.MIDDLE)
        add_rect(s, 13.15, y, 11.54, 1.07, fill, line=LIGHT, radius=True)
        add_text(s, right, 13.42, y + 0.10, 10.98, 0.84, size=14, color=DARK, valign=MSO_ANCHOR.MIDDLE)
    add_text(s, "기대하는 활용 가치 (아직 현장 효과로 검증되지 않음)", 1.45, 9.12, 23.35, 0.40, size=17, bold=True, color=GREEN)
    benefits = [
        ("확인 지원", "기록 공백이 큰 권역을 추가 조사 후보로 드러냄"),
        ("정책 비교", "예산·미충족·근거 차이를 담당자가 함께 검토"),
        ("대응·설명", "불참 재계획과 판단 전제를 설명 가능한 초안으로 제공"),
    ]
    for i, (head, body) in enumerate(benefits):
        x = 1.44 + i * 8.12
        add_rect(s, x, 9.67, 7.62, 1.36, MINT, line=MINT, radius=True)
        add_text(s, head, x + 0.22, 9.81, 1.75, 0.34, size=15, bold=True, color=GREEN)
        add_text(s, body, x + 2.05, 9.76, 5.28, 0.92, size=13, color=DARK)
    add_rect(s, 1.43, 11.36, 23.60, 1.05, GOLD_BG, line=GOLD_BG, radius=True)
    add_text(s, "향후 검증", 1.69, 11.49, 1.75, 0.34, size=15, bold=True, color=GOLD)
    add_text(s, "주민·지자체 파일럿, 가중치·단가 합의, 개인정보/보안·조달 적합성, 실제 서비스 접근성·비용 효과를 별도로 검증해야 합니다.",
             3.52, 11.43, 20.95, 0.62, size=14, color=DARK)
    add_text(s, "기대효과는 검증된 편익·절감액이 아닙니다. 최종 결정과 행정 승인은 담당자에게 남습니다.",
             1.50, 12.75, 23.35, 0.35, size=11, color=GRAY)

    prs.save(PPT_OUT)
    return data, ex, artifact_sha


def set_cell_text(cell, text, *, font_size=9.2, bold=False, color="24342D", first_page_break=False):
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    lines = text.split("\n")
    for idx, line in enumerate(lines):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        if idx == 0 and first_page_break:
            p.paragraph_format.page_break_before = True
        run = p.add_run(line)
        run.font.name = "맑은 고딕"
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.color.rgb = RGBColor.from_string(color)
        rpr = run._element.get_or_add_rPr()
        rfonts = rpr.rFonts
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:eastAsia"), "맑은 고딕")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def build_docx(data, ex):
    baseline, balanced = ex["request_count_baseline"], ex["balanced"]
    total_areas = len(ex["request_count_service_units_by_area"])
    doc = Document(str(DOCX_TEMPLATE))
    report = doc.tables[0]
    # The official DOCX bundle includes a second participant application table.
    # Keep the report table untouched in structure and remove only that separate
    # application table from this report copy.
    if len(doc.tables) > 1:
        app_table = doc.tables[1]
        app_table._element.getparent().remove(app_table._element)

    content = {
        1: f"{TECH_NAME}\n(팀으로 참여하는 경우 대표자 1명만 제출)",
        2: "(AI 라이프 아이디어 챌린지를 참고했다면 해당하는 아이디어 번호와 아이디어 명을 기재)\n"
           "예시> 00번. 00000000000000000000000\n"
           f"작성 아이디어명: {IDEA_NAME}\n"
           "아이디어 번호·기존 아이디어명·선정 상태: 제출 전 사용자 확인 필요",
        3: "요청 기록이 적은 마을을 필요가 없는 곳으로 단정하지 않고, 근거가 부족하거나 충돌하면 조사 필요로 남긴다.\n"
           "AI 보조 구조화와 OR-Tools CP-SAT으로 네 정책 공급안을 비교하고 최종 선택은 담당자가 검토한다.\n"
           "현재 결과는 시제품·합성자료 기준이며 주민 데이터와 현장 효과는 검증 전이다.",
        4: "근거 공백을 드러내면서 예산·공급·이동 제약 아래 생활서비스 공급안을 비교하는 시제품 구현.\n"
           "정책 선택과 승인은 담당자에게 남기고, 현장 적용 전 운영값·가중치·보안 절차를 검증한다.",
        5: "ㅇ 구현: 근거 부족·충돌을 조사 필요/UNKNOWN으로 표시하고 담당자 검토 전 승인 근거로 사용하지 않는다.\n"
           "ㅇ AI 경로: Gemini 설정 시 구조화 출력, 미설정 시 결정적 fallback을 쓴다. 이번 로컬 E2E는 외부 모델 키 없이 수행했다.\n"
           "ㅇ 최적화: OR-Tools CP-SAT으로 효율 중심·균형·최소보장·소외 최소화 4개 정책을 계산한다.\n"
           "ㅇ 시연 검증: 격리된 로컬 Playwright E2E 1건에서 네 정책 비교, 합성 제공자 불참 후 재계획, 검토자료 PDF/CSV 흐름 통과.\n"
           f"ㅇ 별도 통제 실험({SIM_LABEL}): {data['region']} {total_areas}개 권역·예산 {ex['budget_won']:,}원·seed {data['seed']}·합성 조사 필요 {baseline['survey_required_areas']}개. 요청량 기준안/균형안은 각 {baseline['served_units']}단위, 배정 권역 {baseline['covered_areas']}/{total_areas} 대 {balanced['covered_areas']}/{total_areas}, 조사 필요 배정 {baseline['survey_required_covered']}/{baseline['survey_required_areas']} 대 {balanced['survey_required_covered']}/{balanced['survey_required_areas']}. 실제 주민 서비스·효과가 아니다.\n"
           "※ 개발 기술, 제품에 대한 기본 구현 내용, 탐구 내용 및 방법 등",
        6: "ㅇ 현장 주민자료와 제공자 가격·가용량이 부족해 공개·합성·미확인 입력을 구분하고 조사 필요 상태를 유지했다.\n"
           "ㅇ 이동·일정 제약과 solver 시간/최적성 한계를 고려해 결과 상태와 가정을 표시했다.\n"
           "ㅇ 개인정보 및 기관 운영 보안·행정 적합성은 실제 공공기관 파일럿에서 추가 검증해야 한다.",
        7: "ㅇ 기록 공백이 큰 지역과 정책별 미충족 결과를 함께 보여 담당자의 추가 조사·검토를 지원할 수 있다.\n"
           "ㅇ 공급자 불참 뒤 대체 계획을 비교하고 판단 전제·제약을 설명하는 데 도움을 줄 수 있다.\n"
           "ㅇ 서비스 접근성 개선·비용 절감은 아직 검증되지 않았으며 현장 평가가 필요하다.",
    }
    # Keep official title row and all seven labels verbatim.
    for row_idx in range(1, 8):
        cell = report.rows[row_idx].cells[1]
        set_cell_text(cell, content[row_idx], font_size=8.7 if row_idx >= 5 else 9.1,
                      first_page_break=(row_idx == 5))
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if row_idx == 2:
            set_cell_shading(cell, "FFF8EA")
    # Keep all report rows together where possible, with the first detail section
    # beginning on a fresh page via pageBreakBefore on both cells in row 5.
    for idx, row in enumerate(report.rows):
        tr_pr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if idx == 5:
            for cell in row.cells:
                cell.paragraphs[0].paragraph_format.page_break_before = True
    doc.core_properties.title = "2026년 AI 라이프 솔루션 챌린지 VillageCoverage 성과보고서"
    doc.core_properties.subject = "공식 Word 성과보고서 양식의 2쪽 이내 작성 대기본"
    doc.core_properties.author = "VillageCoverage"
    doc.save(DOCX_OUT)


def load_font(size, bold=False):
    # Available Korean system typeface on this authoring host. Fonts are not
    # embedded in PPTX/DOCX; Office may substitute if Pretendard/맑은 고딕 is absent.
    path = "/System/Library/Fonts/AppleSDGothicNeo.ttc"
    return ImageFont.truetype(path, size, index=0)


def wrap_pixels(draw, text, font, max_width):
    out, line = [], ""
    for ch in text:
        candidate = line + ch
        if line and draw.textlength(candidate, font=font) > max_width:
            out.append(line)
            line = ch
        else:
            line = candidate
    if line:
        out.append(line)
    return out


def draw_wrapped(draw, xy, text, font, fill, max_width, line_gap=8, max_lines=None):
    x, y = xy
    lines = []
    for para in text.split("\n"):
        if not para:
            lines.append("")
        else:
            lines.extend(wrap_pixels(draw, para, font, max_width))
    if max_lines is not None:
        lines = lines[:max_lines]
    line_height = font.size + line_gap
    for line in lines:
        draw.text((x, y), line, font=font, fill=fill)
        y += line_height
    return y


def render_preview_pages(data, ex, artifact_sha):
    # 16:9 page previews for design review. These are independent raster
    # renderings from the same content/evidence, not Office-engine exports.
    W, H = 1920, 1080
    bg, dark, green, gray, light, mint, gold, goldbg = "#FFFFFF", "#18382D", "#176B4A", "#52615B", "#D8E2DC", "#EAF4EF", "#A66A11", "#FFF5E4"
    titlef, bodyf, smallf = load_font(46), load_font(25), load_font(20)
    headings = ["1. AI 기술명", "2. 아이디어명 / 3. 기술요약", "4. 최종목표 / 5. 개발 내용 및 결과", "6. 개발 과정에서 애로점 / 7. 기대효과"]
    ppt_pages = []
    for i, title in enumerate(headings):
        im = Image.new("RGB", (W,H), bg)
        d = ImageDraw.Draw(im)
        d.rectangle((0,0,W,16), fill=green)
        d.text((90,48), title, font=titlef, fill=dark)
        d.line((90,124,W-90,124), fill=light, width=2)
        if i == 0:
            d.text((95,160), TECH_NAME, font=load_font(30), fill=dark)
            draw_wrapped(d,(95,208),CORE,bodyf,"#52615B",W-190,7,2)
            a=Image.open(EVIDENCE/"ScenarioCompare_local-E2E.png").convert("RGB")
            b=Image.open(EVIDENCE/"Dashboard_map_local-E2E.png").convert("RGB")
            # Keep screenshots uncropped in these preview frames; report screenshots
            # are scaled to fit their own labeled panels.
            panel_y=320; panel_h=585; panel_w=820
            for x, image, label in [(95,a,"네 정책 비교 · 로컬 E2E"),(1005,b,"대시보드 지도 · 로컬 E2E")]:
                d.rounded_rectangle((x,panel_y-48,x+panel_w,panel_y+panel_h+8),radius=20,fill="#FFFFFF",outline=light,width=3)
                d.text((x+20,panel_y-38),label,font=load_font(21),fill=green)
                image.thumbnail((panel_w-36,panel_h-28))
                im.paste(image,(x+18,panel_y+5))
            d.text((95,950),"로컬 Playwright E2E · 합성 데모 데이터 · 2026-10-13 · source b6ff52a · Render 배포 화면 아님",font=smallf,fill=gray)
        elif i == 1:
            d.rounded_rectangle((95,165,W-95,235),radius=18,fill=goldbg)
            d.text((120,184),"이전 아이디어 번호·선정 상태·5점 가점 자격: 사용자 확인 필요",font=bodyf,fill=gold)
            blocks=[("01 문제","요청 기록이 적은 마을을\n필요가 없는 곳으로 오판할 위험"),("02 근거 구조화","AI 보조로 비정형 기록을\n검토 초안으로 정리. 부족·충돌은 조사"),("03 제약 최적화","CP-SAT으로 예산·공급·이동·일정\n조건 아래 네 정책안을 계산"),("04 담당자 검토","담당자가 전제·미배정 사유를\n검토하고 선택. 자동 승인하지 않음")]
            bx, by, bw, bh, gap=95,285,405,235,58
            for n,(h,t) in enumerate(blocks):
                x=bx+n*(bw+gap)
                d.rounded_rectangle((x,by,x+bw,by+bh),radius=20,fill="#F5F8F6",outline=light,width=3)
                d.text((x+22,by+20),h,font=load_font(24),fill=green)
                draw_wrapped(d,(x+22,by+72),t.replace("\n"," "),bodyf,dark,bw-44,8,5)
            d.text((95,590),"시제품의 네 정책",font=load_font(32),fill=dark)
            for n,(h,t) in enumerate([("효율 중심","같은 예산에서 서비스 회차 효율"),("균형","회차·권역 분산·조사 필요 고려"),("소외 최소화","서비스 공백이 긴 권역 우선"),("최소보장","권역별 최소 회차 조건 우선")]):
                x=95+n*440
                d.rounded_rectangle((x,655,x+405,810),radius=16,fill=mint,outline=light,width=2)
                d.text((x+20,675),h,font=load_font(25),fill=green)
                draw_wrapped(d,(x+20,718),t,smallf,dark,365,6,3)
            d.rounded_rectangle((95,855,W-95,963),radius=15,fill="#F5F8F6")
            draw_wrapped(d,(120,880),"AI는 근거를 정리하는 보조 경로, CP-SAT은 명시된 목적·제약을 계산하는 solver입니다. 가중치·단가는 시제품 선택이며 현장 합의값이 아닙니다.",smallf,dark,W-245,7,3)
        elif i == 2:
            d.text((95,160),"구현 범위와 통제 합성 실험",font=load_font(32),fill=dark)
            left=["근거 부족·충돌은 조사 필요/UNKNOWN으로 남김","Gemini 설정 시 구조화 출력, 미설정 시 fallback","CP-SAT 네 정책 계산과 solver 상태 구분","합성 불참 뒤 계획 버전·Memo/PDF/CSV E2E 통과"]
            for j,t in enumerate(left):
                y=235+j*130
                d.rounded_rectangle((95,y,820,y+102),radius=15,fill="#F5F8F6",outline=light,width=2)
                draw_wrapped(d,(120,y+22),t,bodyf,dark,660,8,3)
            tx,ty=865,220
            d.rounded_rectangle((tx,ty,W-95,ty+550),radius=20,fill="#FFFFFF",outline=light,width=3)
            d.text((tx+28,ty+24),"SIMULATED FOR PRE-R&D · 요청량 기준 vs 균형",font=load_font(25),fill=gold)
            baseline, balanced = ex["request_count_baseline"], ex["balanced"]
            total_areas = len(ex["request_count_service_units_by_area"])
            d.text((tx+28,ty+70),f"{data['region']} · {total_areas}개 권역 · 예산 ₩{ex['budget_won']:,} · seed {data['seed']}",font=smallf,fill=gray)
            vals=[("서비스 단위",str(baseline["served_units"]),str(balanced["served_units"])),("배정 권역",f"{baseline['covered_areas']} / {total_areas}",f"{balanced['covered_areas']} / {total_areas}"),("미배정 권역",str(total_areas-baseline["covered_areas"]),str(total_areas-balanced["covered_areas"])),("조사 필요 권역 배정",f"{baseline['survey_required_covered']} / {baseline['survey_required_areas']}",f"{balanced['survey_required_covered']} / {balanced['survey_required_areas']}"),("모델 이동비",f"₩{baseline['travel_cost_won']:,}",f"₩{balanced['travel_cost_won']:,}")]
            d.rectangle((tx+28,ty+125,W-123,ty+176),fill=mint)
            d.text((tx+40,ty+137),"지표",font=smallf,fill=green); d.text((tx+420,ty+137),"요청량 기준",font=smallf,fill=green); d.text((tx+675,ty+137),"균형",font=smallf,fill=green)
            for j,(a,b,c) in enumerate(vals):
                yy=ty+176+j*66
                d.line((tx+28,yy+64,W-123,yy+64),fill=light,width=2)
                d.text((tx+40,yy+14),a,font=smallf,fill=dark); d.text((tx+460,yy+14),b,font=smallf,fill=dark); d.text((tx+700,yy+14),c,font=smallf,fill=dark)
            draw_wrapped(d,(tx+28,ty+512),"합성 관측 4건 미만=조사 필요. 이동비는 모델 추정치로 실제 집행비가 아닙니다.",smallf,gray,850,5,2)
            d.rounded_rectangle((95,805,W-95,965),radius=18,fill=goldbg)
            draw_wrapped(d,(125,830),"현장 서비스·효과를 입증하지 않습니다. 이 실험은 두 정책만 비교하며 UI의 네 정책 E2E와 별도입니다.\n실험 JSON SHA-256: "+artifact_sha,smallf,dark,W-250,8,4)
        else:
            cols=[(95,"개발 애로점",["현장 주민 자료 부족","공급자 가격·용량 미확정","이동·일정·solver 복잡도","데이터 출처·개인정보·보안"]),(980,"대응 설계",["조사 필요/UNKNOWN 유지","합성 운영값 분리·불참 재계산","solver 상태·가정 함께 표시","출처 구분·담당자 검토 유지"])]
            for x,h,items in cols:
                d.text((x,160),h,font=load_font(32),fill=dark)
                for j,t in enumerate(items):
                    y=220+j*118
                    d.rounded_rectangle((x,y,x+820,y+96),radius=15,fill="#F5F8F6",outline=light,width=2)
                    draw_wrapped(d,(x+22,y+22),t,bodyf,dark,770,7,2)
            d.text((95,735),"기대하는 활용 가치 · 현장 효과로 검증 전",font=load_font(30),fill=green)
            for j,(h,t) in enumerate([("확인 지원","기록 공백이 큰 권역을 조사 후보로 드러냄"),("정책 비교","예산·미충족·근거 차이를 검토"),("대응·설명","불참 재계획과 판단 전제를 초안으로 제공")]):
                x=95+j*590
                d.rounded_rectangle((x,795,x+545,900),radius=14,fill=mint)
                d.text((x+18,810),h,font=load_font(23),fill=green)
                draw_wrapped(d,(x+160,808),t,smallf,dark,360,5,3)
            d.rounded_rectangle((95,930,W-95,1025),radius=14,fill=goldbg)
            draw_wrapped(d,(120,948),"실제 접근성·비용 효과는 미입증. 주민·지자체 파일럿, 가중치·단가 합의, 보안·조달 적합성 검증이 필요합니다.",smallf,dark,W-245,6,3)
        d.line((90,1038,W-90,1038),fill=light,width=2)
        d.text((95,1044),"VillageCoverage · Pre-R&D 시제품 · 최종 정책 선택과 승인은 담당자",font=load_font(17),fill=gray)
        path=PREVIEW/f"ppt-slide-{i+1:02}.png"
        im.save(path)
        ppt_pages.append(im)
    ppt_pages[0].save(PPT_PDF, "PDF", resolution=144.0, save_all=True, append_images=ppt_pages[1:])

    # Two-page Word proof view mirrors the specified official field order and
    # page split. It does not assert Word's own pagination.
    page_w,page_h=1240,1754
    docs=[]
    page1=Image.new("RGB",(page_w,page_h),"white"); d=ImageDraw.Draw(page1)
    title=load_font(34); lab=load_font(20); text=load_font(19); foot=load_font(15)
    d.text((70,55),"2026년 AI 라이프 솔루션 챌린지 성과보고서",font=title,fill="#18382D")
    d.text((70,103),"공식 Word 양식 · 최대 2쪽 작성 대기본",font=foot,fill="#52615B")
    d.line((70,145,page_w-70,145),fill="#176B4A",width=4)
    fields=[("AI 기술명",TECH_NAME+"\n(팀으로 참여하는 경우 대표자 1명만 제출)"),("아이디어명",IDEA_NAME+"\n(AI 라이프 아이디어 챌린지를 참고했다면 해당하는 아이디어 번호와 아이디어 명을 기재)\n예시> 00번. 00000000000000000000000\n아이디어 번호·선정 상태: 제출 전 사용자 확인 필요"),("기술요약 (3줄 이내)","요청 기록이 적은 마을을 필요 없는 곳으로 단정하지 않고, 근거 부족·충돌은 조사 필요로 남긴다. AI 보조 구조화와 OR-Tools CP-SAT으로 네 정책안을 비교하고 최종 선택은 담당자가 검토한다. 시제품·합성자료 기준이며 주민 데이터와 현장 효과는 검증 전이다."),("최종목표 (2줄 이내)","근거 공백을 드러내면서 예산·공급·이동 제약 아래 생활서비스 공급안을 비교하는 시제품 구현. 정책 선택·승인은 담당자에게 남기고, 현장 적용 전 운영값·가중치·보안 절차를 검증한다.")]
    y=175
    for head,body in fields:
        d.rounded_rectangle((70,y,page_w-70,y+47),radius=8,fill="#EAF4EF")
        d.text((88,y+10),head,font=lab,fill="#176B4A")
        y+=61
        y=draw_wrapped(d,(95,y),body,text,"#24342D",page_w-190,8,5)+20
    d.text((70,page_h-70),"양식 항목명/순서 유지 · PII/동의/서명 없음 · 워드 렌더링은 별도 확인 필요",font=foot,fill="#52615B")
    docs.append(page1)

    page2=Image.new("RGB",(page_w,page_h),"white"); d=ImageDraw.Draw(page2)
    d.text((70,55),"성과보고서 · 계속",font=title,fill="#18382D")
    d.line((70,120,page_w-70,120),fill="#176B4A",width=4)
    y=150
    items=[
        ("5. 개발 내용 및 결과",f"• 근거 부족·충돌은 조사 필요/UNKNOWN으로 표시; 담당자 검토 전 승인 근거로 사용하지 않음.\n• Gemini 설정 시 구조화 출력, 미설정 시 결정적 fallback. 이번 E2E는 외부 모델 키 없이 실행.\n• CP-SAT 네 정책 계산; 합성 제공자 불참 후 재계획, 검토자료 PDF/CSV를 로컬 Playwright E2E 1건에서 통과.\n• SIMULATED FOR PRE-R&D: {data['region']} {total_areas}권역·₩{ex['budget_won']:,}·seed {data['seed']}. 요청량 기준/균형 각 {baseline['served_units']}단위; 배정 {baseline['covered_areas']}/{total_areas} 대 {balanced['covered_areas']}/{total_areas}; 조사 필요 배정 {baseline['survey_required_covered']}/{baseline['survey_required_areas']} 대 {balanced['survey_required_covered']}/{balanced['survey_required_areas']}. 실제 서비스·효과 아님."),
        ("6. 개발 과정에서 애로점","• 현장 주민자료와 제공자 가격·가용량 부족: 공개·합성·미확인 입력 분리, 조사 필요 유지.\n• 이동·일정·solver 제약: 결과 상태와 가정을 표시하고 최적성 미증명은 구분.\n• 개인정보·기관 보안·행정 적합성은 실제 공공기관 파일럿 검증 전."),
        ("7. 기대효과","• 기록 공백이 큰 권역과 정책별 미충족 결과를 함께 보여 추가 조사·검토를 지원할 수 있다.\n• 공급자 불참 후 계획을 비교하고 판단 전제·제약을 설명하는 데 도움을 줄 수 있다.\n• 접근성 개선·비용 절감은 미입증이며 현장 평가가 필요하다.")]
    for head,body in items:
        d.rounded_rectangle((70,y,page_w-70,y+47),radius=8,fill="#EAF4EF")
        d.text((88,y+10),head,font=lab,fill="#176B4A")
        y+=63
        y=draw_wrapped(d,(95,y),body,text,"#24342D",page_w-190,7,16)+26
    d.text((70,page_h-70),"검증 미리보기 · Microsoft Word/LibreOffice PDF 변환본 아님",font=foot,fill="#A66A11")
    docs.append(page2)
    docs[0].save(DOCX_PDF,"PDF",resolution=150.0,save_all=True,append_images=docs[1:])
    for i,im in enumerate(docs,1):
        im.save(PREVIEW/f"docx-page-{i:02}.png")


def main():
    prep_evidence()
    data, ex, artifact_sha = build_ppt()
    build_docx(data, ex)
    render_preview_pages(data, ex, artifact_sha)
    print(f"PPTX={PPT_OUT} bytes={PPT_OUT.stat().st_size}")
    print(f"DOCX={DOCX_OUT} bytes={DOCX_OUT.stat().st_size}")
    print(f"PPTX_PREVIEW_PDF={PPT_PDF} bytes={PPT_PDF.stat().st_size}")
    print(f"DOCX_PREVIEW_PDF={DOCX_PDF} bytes={DOCX_PDF.stat().st_size}")
    print(f"EXPERIMENT_SHA256={artifact_sha}")


if __name__ == "__main__":
    main()
